"""
Homework 3 — Fine-tuning Benchmark
===================================
Compares google/gemma-4-E4B-it (base) vs the LoRA-adapted version
(habert75/autogen-gemma4-4bit-lora-adapter) trained on hotel-policies data.

Two question sets:
  1. Hotel-policy questions  → fine-tuned model expected to answer better.
  2. General / moral questions → both models should perform similarly (sanity check).

Output: benchmark_results.xlsx  (formatted Excel table, like Homework 1)

Usage (RunPod or any machine with GPU):
    pip install transformers peft bitsandbytes accelerate torch openpyxl
    python benchmark.py
"""

import gc
import os
import re
import torch
import openpyxl
from datetime import datetime
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from transformers import AutoModelForCausalLM, AutoTokenizer, BitsAndBytesConfig
from peft import PeftModel
from huggingface_hub import login, HfApi, create_repo

# ── Config ────────────────────────────────────────────────────────────────────
# Token read from environment — never hardcode in committed files.
# On RunPod: export HF_TOKEN='hf_...' in terminal before running.
HF_TOKEN       = os.environ.get("HF_TOKEN", "")
if not HF_TOKEN:
    raise ValueError("Set the HF_TOKEN environment variable first.")
HF_USER        = "habert75"
BASE_MODEL     = "google/gemma-4-E4B-it"
ADAPTER_REPO   = "habert75/autogen-gemma4-4bit-lora-adapter"
RESULTS_REPO   = f"{HF_USER}/homework3-benchmark-results"   # dataset repo to push Excel to
OUTPUT_FILE    = "benchmark_results.xlsx"
MAX_NEW_TOKENS = 80

# ── Questions ─────────────────────────────────────────────────────────────────
# Each entry: (question, expected_keywords, category)
# expected_keywords: list of strings that a correct answer should contain (case-insensitive).
# Scoring: 1 point per matching keyword found, normalised to 0-100%.

HOTEL_QUESTIONS = [
    (
        "What time is check-in and check-out?",
        ["2", "14", "12", "check-in", "check-out", "noon", "p.m"],
        "Hotel Policy",
    ),
    (
        "What is the cancellation policy?",
        ["72 hours", "penalty", "prepay", "written", "cancel", "no-show", "no show"],
        "Hotel Policy",
    ),
    (
        "Are pets allowed in the hotel?",
        ["5 kg", "35", "pet", "allowed", "surcharge", "approval"],
        "Hotel Policy",
    ),
    (
        "How can I make a booking at the hotel?",
        ["email", "website", "office@", "online", "reservation", "credit card"],
        "Hotel Policy",
    ),
    (
        "Is breakfast included in the room rate?",
        ["breakfast", "included", "à la carte", "a la carte"],
        "Hotel Policy",
    ),
    (
        "What payment methods does the hotel accept?",
        ["visa", "mastercard", "american express", "maestro", "cash", "bank transfer"],
        "Hotel Policy",
    ),
    (
        "What is the penalty for smoking in a hotel room?",
        ["75", "EUR", "penalty", "forbidden", "smoke"],
        "Hotel Policy",
    ),
    (
        "Can visitors stay in a guest's room overnight?",
        ["11", "p.m", "registration", "reception", "visitor", "identity"],
        "Hotel Policy",
    ),
    (
        "Is early check-in available and what does it cost?",
        ["10", "a.m", "25%", "early", "additional", "cost"],
        "Hotel Policy",
    ),
    (
        "Is parking available at the hotel?",
        ["free", "parking", "underground", "10 EUR", "territory"],
        "Hotel Policy",
    ),
]

GENERAL_QUESTIONS = [
    (
        "Is it ever morally acceptable to lie to protect someone's feelings?",
        ["compassion", "harm", "truth", "situation", "context", "empathy", "ethical"],
        "General / Moral",
    ),
    (
        "What are three tips for staying productive while working from home?",
        ["schedule", "distraction", "break", "routine", "focus", "workspace", "time"],
        "General / Moral",
    ),
    (
        "Explain briefly what machine learning is.",
        ["data", "model", "learn", "pattern", "algorithm", "train", "predict"],
        "General / Moral",
    ),
    (
        "Should wealthy countries do more to help poorer nations?",
        ["aid", "inequality", "responsibility", "develop", "economic", "support"],
        "General / Moral",
    ),
    (
        "What is the capital of France?",
        ["paris"],
        "General / Moral",
    ),
]

ALL_QUESTIONS = HOTEL_QUESTIONS + GENERAL_QUESTIONS


# ── Model helpers ─────────────────────────────────────────────────────────────

def get_bnb_config():
    return BitsAndBytesConfig(
        load_in_4bit=True,
        bnb_4bit_use_double_quant=True,
        bnb_4bit_compute_dtype=torch.bfloat16,
        bnb_4bit_quant_type="nf4",
    )


def load_base(quantize=True):
    kwargs = dict(token=HF_TOKEN)
    if quantize and torch.cuda.is_available():
        kwargs["quantization_config"] = get_bnb_config()
        kwargs["dtype"] = torch.bfloat16
        kwargs["device_map"] = "auto"
    else:
        kwargs["dtype"] = torch.float32
    print(f"  Loading {BASE_MODEL}…")
    model = AutoModelForCausalLM.from_pretrained(BASE_MODEL, **kwargs)
    tokenizer = AutoTokenizer.from_pretrained(BASE_MODEL, token=HF_TOKEN)
    model.eval()
    return model, tokenizer


def load_finetuned():
    model, tokenizer = load_base(quantize=True)
    tokenizer = AutoTokenizer.from_pretrained(ADAPTER_REPO, token=HF_TOKEN)
    print(f"  Attaching adapter {ADAPTER_REPO}…")
    peft_kwargs = {"token": HF_TOKEN}
    if torch.cuda.is_available():
        peft_kwargs["device_map"] = "auto"
    model = PeftModel.from_pretrained(model, ADAPTER_REPO, **peft_kwargs)
    model.eval()
    return model, tokenizer


INFER_DEVICE = "cuda:0" if torch.cuda.is_available() else "cpu"


def generate(model, tokenizer, prompt, max_new_tokens=MAX_NEW_TOKENS):
    messages = [{"role": "user", "content": prompt}]
    formatted = tokenizer.apply_chat_template(
        messages, tokenize=False, add_generation_prompt=True
    )
    inputs = tokenizer(formatted, return_tensors="pt").to(INFER_DEVICE)
    with torch.no_grad():
        outputs = model.generate(
            **inputs,
            max_new_tokens=max_new_tokens,
            do_sample=False,
            pad_token_id=tokenizer.pad_token_id,
            eos_token_id=tokenizer.eos_token_id,
        )
    generated = outputs[0][inputs["input_ids"].shape[1]:]
    return tokenizer.decode(generated, skip_special_tokens=True).strip()


def score(answer, keywords):
    """Return score 0-100 based on how many keywords appear in the answer."""
    if not keywords:
        return 0
    answer_lower = answer.lower()
    hits = sum(1 for kw in keywords if kw.lower() in answer_lower)
    return round(hits / len(keywords) * 100)


# ── Run benchmark ─────────────────────────────────────────────────────────────

def run_benchmark(model, tokenizer, label):
    print(f"\n{'='*55}")
    print(f"  Running: {label}")
    print(f"{'='*55}")
    results = []
    for i, (question, keywords, category) in enumerate(ALL_QUESTIONS, 1):
        print(f"  [{i}/{len(ALL_QUESTIONS)}] {question[:60]}…")
        answer = generate(model, tokenizer, question)
        s = score(answer, keywords)
        results.append({
            "question":  question,
            "category":  category,
            "keywords":  ", ".join(keywords),
            "answer":    answer,
            "score":     s,
        })
        print(f"         → score {s}%")
    return results


# ── Excel output ──────────────────────────────────────────────────────────────

def hex_fill(hex_color):
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")

COLORS = {
    "header_hotel":   "1B5E20",   # dark green
    "header_general": "1A237E",   # dark blue
    "header_meta":    "37474F",   # slate
    "row_hotel":      "E8F5E9",   # light green
    "row_general":    "E8EAF6",   # light blue
    "row_alt_hotel":  "C8E6C9",
    "row_alt_gen":    "C5CAE9",
    "score_hi":       "66BB6A",   # green  ≥ 70
    "score_mid":      "FFA726",   # orange 40-69
    "score_lo":       "EF5350",   # red    < 40
    "white":          "FFFFFF",
}

thin = Side(style="thin", color="BDBDBD")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def cell_style(ws, row, col, value, bold=False, fill=None, font_color="000000",
               wrap=True, h_align="left", font_size=10):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(bold=bold, color=font_color, size=font_size)
    if fill:
        c.fill = hex_fill(fill)
    c.alignment = Alignment(wrap_text=wrap, vertical="top", horizontal=h_align)
    c.border = border
    return c


def score_color(s):
    if s >= 70:
        return COLORS["score_hi"]
    if s >= 40:
        return COLORS["score_mid"]
    return COLORS["score_lo"]


def write_excel(base_results, ft_results, path):
    wb = openpyxl.Workbook()

    # ── Sheet 1: Side-by-side comparison ──────────────────────────────────────
    ws = wb.active
    ws.title = "Comparison"
    ws.freeze_panes = "A3"

    # Title row
    ws.merge_cells("A1:H1")
    t = ws["A1"]
    t.value = (f"Homework 3 — Fine-tuning Benchmark  |  "
               f"Base: {BASE_MODEL}  vs  Fine-tuned: {ADAPTER_REPO}  |  "
               f"{datetime.now().strftime('%Y-%m-%d %H:%M')}")
    t.font = Font(bold=True, size=12, color="FFFFFF")
    t.fill = hex_fill(COLORS["header_meta"])
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    # Column headers
    headers = ["#", "Category", "Question",
               "Base Model Answer", "Base Score %",
               "Fine-tuned Answer", "FT Score %", "Δ Score"]
    widths   = [4,   14,           42,
                52,               13,
                52,               10,          10]
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        ws.column_dimensions[get_column_letter(col)].width = w
        cell_style(ws, 2, col, h, bold=True, fill=COLORS["header_meta"],
                   font_color="FFFFFF", h_align="center")

    # Data rows
    hotel_count   = len(HOTEL_QUESTIONS)
    for i, (b, f) in enumerate(zip(base_results, ft_results)):
        r = i + 3
        is_hotel = i < hotel_count
        alt = (i % 2 == 1)
        if is_hotel:
            row_fill = COLORS["row_alt_hotel"] if alt else COLORS["row_hotel"]
        else:
            row_fill = COLORS["row_alt_gen"] if alt else COLORS["row_general"]

        delta = f["score"] - b["score"]
        delta_color = COLORS["score_hi"] if delta > 0 else (COLORS["score_lo"] if delta < 0 else "FFFFFF")

        cell_style(ws, r, 1, i + 1,        fill=row_fill, h_align="center")
        cell_style(ws, r, 2, b["category"], fill=row_fill)
        cell_style(ws, r, 3, b["question"], fill=row_fill)
        cell_style(ws, r, 4, b["answer"],   fill=row_fill)
        cell_style(ws, r, 5, b["score"],    fill=score_color(b["score"]), h_align="center", bold=True, font_color="FFFFFF")
        cell_style(ws, r, 6, f["answer"],   fill=row_fill)
        cell_style(ws, r, 7, f["score"],    fill=score_color(f["score"]), h_align="center", bold=True, font_color="FFFFFF")
        cell_style(ws, r, 8, f"{delta:+d}", fill=delta_color, h_align="center", bold=True, font_color="FFFFFF")
        ws.row_dimensions[r].height = 80

    # ── Sheet 2: Summary statistics ───────────────────────────────────────────
    ws2 = wb.create_sheet("Summary")

    def avg(results, category=None):
        subset = [r["score"] for r in results if category is None or r["category"] == category]
        return round(sum(subset) / len(subset), 1) if subset else 0

    rows_summary = [
        ("Metric", "Base Model", "Fine-tuned", "Δ"),
        ("Overall avg score (%)",
         avg(base_results), avg(ft_results),
         round(avg(ft_results) - avg(base_results), 1)),
        ("Hotel Policy avg score (%)",
         avg(base_results, "Hotel Policy"), avg(ft_results, "Hotel Policy"),
         round(avg(ft_results, "Hotel Policy") - avg(base_results, "Hotel Policy"), 1)),
        ("General / Moral avg score (%)",
         avg(base_results, "General / Moral"), avg(ft_results, "General / Moral"),
         round(avg(ft_results, "General / Moral") - avg(base_results, "General / Moral"), 1)),
        ("Questions answered (total)", len(base_results), len(ft_results), "—"),
        ("Base model", BASE_MODEL, "", ""),
        ("Fine-tuned adapter", ADAPTER_REPO, "", ""),
        ("Max new tokens", MAX_NEW_TOKENS, MAX_NEW_TOKENS, ""),
        ("Decoding", "Greedy (do_sample=False)", "Greedy (do_sample=False)", ""),
        ("Run date", datetime.now().strftime("%Y-%m-%d %H:%M"), "", ""),
    ]

    ws2.column_dimensions["A"].width = 35
    ws2.column_dimensions["B"].width = 55
    ws2.column_dimensions["C"].width = 55
    ws2.column_dimensions["D"].width = 12

    for r, row_data in enumerate(rows_summary, 1):
        is_header = r == 1
        for c, val in enumerate(row_data, 1):
            fill = COLORS["header_meta"] if is_header else (COLORS["row_hotel"] if r % 2 == 0 else "FFFFFF")
            fc = "FFFFFF" if is_header else "000000"
            cell_style(ws2, r, c, val, bold=is_header, fill=fill, font_color=fc, h_align="center" if c > 1 else "left")
        ws2.row_dimensions[r].height = 22

    wb.save(path)
    print(f"\n✅  Results saved to: {path}")


# ── Push Excel to Hugging Face Hub ────────────────────────────────────────────

def push_to_hub(local_path, repo_id, token):
    """Upload the Excel file to a HuggingFace dataset repo."""
    print(f"\nPushing {local_path} to HuggingFace Hub ({repo_id})…")
    create_repo(repo_id, repo_type="dataset", token=token, exist_ok=True)
    api = HfApi()
    api.upload_file(
        path_or_fileobj=local_path,
        path_in_repo=local_path,
        repo_id=repo_id,
        repo_type="dataset",
        token=token,
        commit_message=f"Add benchmark results — {datetime.now().strftime('%Y-%m-%d %H:%M')}",
    )
    print(f"✅  Uploaded to: https://huggingface.co/datasets/{repo_id}")


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    login(token=HF_TOKEN)
    device_label = "CUDA" if torch.cuda.is_available() else "CPU (slow)"
    print(f"\nDevice: {device_label}")

    # 1. Fine-tuned model first (adapter already downloaded)
    print("\n[1/2] Loading FINE-TUNED model…")
    ft_model, ft_tok = load_finetuned()
    ft_results = run_benchmark(ft_model, ft_tok, "Fine-tuned (base + adapter)")

    # Free VRAM before loading the plain base
    del ft_model
    gc.collect()
    if torch.cuda.is_available():
        torch.cuda.empty_cache()

    # 2. Base model (no adapter)
    print("\n[2/2] Loading BASE model…")
    base_model, base_tok = load_base(quantize=True)
    base_results = run_benchmark(base_model, base_tok, "Base model (no adapter)")

    del base_model
    gc.collect()
    if torch.cuda.is_available():
        torch.cuda.empty_cache()

    # 3. Write Excel
    print("\nWriting Excel report…")
    write_excel(base_results, ft_results, OUTPUT_FILE)

    # 4. Push Excel to Hugging Face Hub
    push_to_hub(OUTPUT_FILE, RESULTS_REPO, HF_TOKEN)

    # 5. Print quick summary to console
    print("\n" + "=" * 55)
    print("  SUMMARY")
    print("=" * 55)
    categories = ["Hotel Policy", "General / Moral"]
    for cat in categories:
        b_avg = sum(r["score"] for r in base_results if r["category"] == cat)
        f_avg = sum(r["score"] for r in ft_results  if r["category"] == cat)
        n = sum(1 for r in base_results if r["category"] == cat)
        b_avg /= n; f_avg /= n
        print(f"  {cat:<22} Base={b_avg:.0f}%  FT={f_avg:.0f}%  Δ={f_avg-b_avg:+.0f}%")
    print("=" * 55)


if __name__ == "__main__":
    main()
